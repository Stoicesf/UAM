# -*- coding: utf-8 -*-
"""Build 10-slide advisor PPT for AC-DSGF (CN)."""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "paper" / "ac_dsgf_cn" / "figures"
OUT = ROOT / "paper" / "ac_dsgf_cn" / "ppt" / "AC_DSGF_导师汇报.pptx"

# Colors
INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x0B, 0x6E, 0x4F)
MUTED = RGBColor(0x5C, 0x6B, 0x73)


def _set_run(run, size=18, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Microsoft YaHei"


def add_title_bar(slide, title: str):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.9))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    _set_run(run, 22, True, RGBColor(0xFF, 0xFF, 0xFF))


def add_bullets(slide, lines, left=0.7, top=1.2, width=12, height=5.5, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        run = p.add_run()
        run.text = line
        _set_run(run, size, False, INK)
        p.space_after = Pt(10)


def blank():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    # use blank layout
    layout = prs.slide_layouts[6]
    return prs, layout


def main():
    prs, layout = blank()

    # 1 Title
    s = prs.slides.add_slide(layout)
    add_title_bar(s, "AC-DSGF 导师汇报")
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(3))
    tf = box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "面向通信受限无人机蜂群协同导航的\n自适应通信拓扑优化方法"
    _set_run(r, 32, True, INK)
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "\n通信效率优化型协同智能 · RA-L 投稿配套中文版\n主张：可比任务性能 + 通信约降低两个数量级"
    _set_run(r2, 16, False, MUTED)

    # 2 Background
    s = prs.slides.add_slide(layout)
    add_title_bar(s, "研究背景")
    add_bullets(s, [
        "• UAV 数量增加 → 协同对通信依赖增强",
        "• 固定 / 稠密通信开销随规模近似平方或高度增长",
        "• 固定拓扑无法随任务态势自适应“何时、与谁、通信多少”",
        "• 工程问题：有限带宽 / 能耗 / 时延下如何保持有效协同？",
    ])

    # 3 Prior problems
    s = prs.slides.add_slide(layout)
    add_title_bar(s, "已有方法问题")
    add_bullets(s, [
        "• Full Communication：通信爆炸，难部署",
        "• Random Drop：可能损害关键协作链路",
        "• Fixed Graph（半径 / 稠密）：无法任务驱动适应",
        "• 注意力 ≠ 通信决策：仍在预定义支撑上聚合",
    ], size=20)

    # 4 Core idea
    s = prs.slides.add_slide(layout)
    add_title_bar(s, "核心思想")
    box = s.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(2.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "让无人机自主决定：\n“什么时候通信、和谁通信、通信多少。”"
    _set_run(r, 28, True, ACCENT)
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = "\n通信拓扑 = 策略优化中的可学习决策变量（非训练后随机剪枝）"
    _set_run(r2, 16, False, MUTED)

    # 5 Framework
    s = prs.slides.add_slide(layout)
    add_title_bar(s, "算法框架（Fig.1）")
    img = FIG / "Fig1_framework.png"
    if img.exists():
        s.shapes.add_picture(str(img), Inches(1.2), Inches(1.2), width=Inches(10.8))

    # 6 Contributions
    s = prs.slides.add_slide(layout)
    add_title_bar(s, "创新点（三条）")
    add_bullets(s, [
        "C1  可学习通信拓扑：g_ij 由任务状态驱动",
        "C2  通信预算联合优化：J = E[R] − λ C，Top-K 约束",
        "C3  Residual Guidance：a = π(o) + βΔ(Φ)，稳定稀疏协同",
        "",
        "非声称：Success 超越 · causal utility · AC-DSGF++",
    ], size=20)

    # 7 Setup
    s = prs.slides.add_slide(layout)
    add_title_bar(s, "实验环境")
    add_bullets(s, [
        "• 仿真：VMAS continuous navigation",
        "• 训练：MAPPO（CTDE）",
        "• 主表：16 UAV × 5 seeds × 200-ep eval",
        "• 对比：MAPPO / GAT / Transformer / DSGF / AC-DSGF",
        "• 指标：Success · Comm（软质量）· CEI",
    ])

    # 8 Core results
    s = prs.slides.add_slide(layout)
    add_title_bar(s, "核心结果（Table I + Fig.4）")
    add_bullets(s, [
        "• DSGF：Success 4.01% · Comm 39.27 · CEI 0.001",
        "• AC-DSGF：Success 3.95% · Comm 0.43 · CEI 0.091",
        "• 不是 Success 最高，而是 Success≈ 且 Comm↓≈两个数量级、CEI↑",
        "• 预算 / 静默 / 丢包：诊断趋势，支持 graceful degradation",
    ], top=1.15, height=2.8, size=17)
    img = FIG / "Fig4_pareto.png"
    if img.exists():
        s.shapes.add_picture(str(img), Inches(3.5), Inches(4.0), width=Inches(6.2))

    # 9 Paper contributions mapping
    s = prs.slides.add_slide(layout)
    add_title_bar(s, "论文贡献（对应 RA-L）")
    add_bullets(s, [
        "• Method：Gate + Budget + Residual 闭环",
        "• Analysis：复杂度 O(NKd+Nd)；行为相关 nn_risk≈0.84",
        "• Validation：主表 + 预算 + 静默 + 丢包 + Runtime",
        "• 定位：通信效率优化型机器人系统方法（适合 RA-L）",
    ], size=20)

    # 10 Future
    s = prs.slides.add_slide(layout)
    add_title_bar(s, "下一步（Future Work）")
    add_bullets(s, [
        "• ROS2 / Gazebo / PX4 部署与硬件验证",
        "• 软通信质量 → 真实无线帧 / 能耗映射",
        "• Communication–Motion Co-design（动作—通信协同）",
        "",
        "不回头扩展 AC-DSGF++；v1 冻结投稿，审稿后再定 2.0",
    ], size=20)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print("saved", OUT)


if __name__ == "__main__":
    main()
