# -*- coding: utf-8 -*-
"""Export advisor-readable DOCX from Chinese MD chapters (simplified)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1] / "paper" / "ac_dsgf_cn"
FIG = ROOT / "figures"
OUT = ROOT / "AC_DSGF_CN_导师版.docx"

CHAPTERS = [
    "01_摘要.md",
    "02_引言.md",
    "03_相关工作.md",
    "04_问题定义.md",
    "05_方法.md",
    "06_实验.md",
    "07_复杂度分析.md",
    "08_讨论.md",
    "09_结论.md",
]

IMG_MAP = {
    "Fig1_framework.png": "figures/Fig1_framework.png",
    "Fig2_motivation.png": "figures/Fig2_motivation.png",
    "Fig3_algorithm_flow.png": "figures/Fig3_algorithm_flow.png",
    "Fig4_pareto.png": "figures/Fig4_pareto.png",
    "Fig5_behavior.png": "figures/Fig5_behavior.png",
    "Fig6_runtime.png": "figures/Fig6_runtime.png",
}


def add_md_ish(doc: Document, text: str):
    """Very light MD → paragraphs; insert known figures."""
    lines = text.splitlines()
    buf = []

    def flush():
        nonlocal buf
        if not buf:
            return
        para = doc.add_paragraph("\n".join(buf))
        para.paragraph_format.space_after = Pt(6)
        for run in para.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
        buf = []

    for line in lines:
        s = line.strip()
        if s.startswith("![") and "](" in s:
            flush()
            # extract path
            path = s.split("](")[1].rstrip(")")
            name = Path(path).name
            fp = FIG / name
            if fp.exists():
                doc.add_picture(str(fp), width=Inches(5.8))
                last = doc.paragraphs[-1]
                last.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_paragraph(f"[缺失图片: {name}]")
            continue
        if s.startswith("# "):
            flush()
            doc.add_heading(s[2:], level=1)
            continue
        if s.startswith("## "):
            flush()
            doc.add_heading(s[3:], level=2)
            continue
        if s.startswith("### "):
            flush()
            doc.add_heading(s[4:], level=3)
            continue
        if s.startswith("> "):
            flush()
            p = doc.add_paragraph(s[2:])
            p.runs[0].italic = True if p.runs else None
            continue
        if s.startswith("---"):
            flush()
            continue
        if s == "":
            flush()
            continue
        buf.append(line)
    flush()


def main():
    doc = Document()
    title = doc.add_heading("面向通信受限无人机蜂群协同的自适应通信拓扑学习方法", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(
        "AC-DSGF 中文导师审阅版 · 对应 RA-L 英文冻结稿\n"
        "主张：可比任务性能 + 通信约降低两个数量级 · 不声称 Success 超越"
    )
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for name in CHAPTERS:
        path = ROOT / name
        if not path.exists():
            continue
        doc.add_page_break()
        add_md_ish(doc, path.read_text(encoding="utf-8"))

    # appendix note
    doc.add_page_break()
    doc.add_heading("附录：公式编号 (1)–(6)", level=1)
    add_md_ish(doc, (ROOT / "EQUATIONS.md").read_text(encoding="utf-8"))

    doc.save(str(OUT))
    print("saved", OUT)


if __name__ == "__main__":
    main()
